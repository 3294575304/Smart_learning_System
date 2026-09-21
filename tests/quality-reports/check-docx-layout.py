"""Inspect the actual production DOCX and raster chart coordinates."""
import sys
import zipfile
from io import BytesIO

from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from PIL import Image

output, template, mode = sys.argv[1:]
doc, base = Document(output), Document(template)
assert len(doc.sections) == 2
assert len(doc.tables) == 8
assert len(doc.inline_shapes) == 6
text = "".join(doc.element.body.itertext())
assert "OBJ" not in text.upper()
assert "本次纳入" not in text and "特殊状态或缺失数据" not in text
assert "同口径" not in text and "不相互替代" not in text
assert "有效学生" not in text
goal_figures = [p for p in doc.tables[5].cell(5, 0).paragraphs if p._p.xpath(".//w:drawing") and p.text.startswith("课程目标")]
assert len(goal_figures) == 3
assert all(p.paragraph_format.keep_together is True for p in goal_figures)
assert "课程目标1" in text and "课程目标2" in text and "课程目标3" in text
assert [doc.tables[5].cell(5, 0).tables[0].cell(i, 0).text for i in range(4, 7)] == ["1", "2", "3"]
assert "".join(doc.tables[7]._tbl.itertext()) == "".join(base.tables[7]._tbl.itertext())
for actual, reference in zip(doc.sections, base.sections):
    for prop in ("page_width", "page_height", "left_margin", "right_margin", "top_margin", "bottom_margin"):
        assert getattr(actual, prop) == getattr(reference, prop)
for ti, ri in ((5, 1), (5, 3), (5, 5), (6, 1), (6, 5)):
    assert doc.tables[ti].rows[ri].height is None
    for run in doc.tables[ti].cell(ri, 0)._tc.iter(qn("w:r")):
        if not run.findall(qn("w:t")):
            continue
        fonts = run.find("./" + qn("w:rPr") + "/" + qn("w:rFonts"))
        assert fonts.get(qn("w:ascii")) == "Times New Roman"
        assert fonts.get(qn("w:eastAsia")) == "宋体"
        assert not any("Theme" in key for key in fonts.attrib)
        assert run.find("./" + qn("w:rPr") + "/" + qn("w:sz")).get(qn("w:val")) in {"18", "21", "24"}
for shape in list(doc.inline_shapes)[2:]:
    assert 0.55 < shape.height / shape.width < 0.7
    parent = shape._inline.getparent().getparent().getparent()
    assert "图" in "".join(parent.itertext())  # Picture and caption cannot paginate separately.
with zipfile.ZipFile(output) as actual, zipfile.ZipFile(template) as reference:
    editable = {"[Content_Types].xml", "docProps/core.xml", "word/_rels/document.xml.rels", "word/document.xml", "word/settings.xml", "word/fontTable.xml"}
    for name in reference.namelist():
        if not name.endswith("/") and name not in editable:
            assert actual.read(name) == reference.read(name), name
    fonts = etree.fromstring(actual.read("word/fontTable.xml"))
    title_font = next(font for font in fonts if font.get(qn("w:name")) == "方正小标宋简体")
    assert title_font.find(qn("w:altName")).get(qn("w:val")) == "微软雅黑"
    assert title_font.find(qn("w:embedRegular")) is None
    charts = [Image.open(BytesIO(actual.read(name))).convert("RGB") for name in actual.namelist() if name.startswith("word/media/") and name.endswith(".png")]
    scatter = [chart for chart in charts if chart.size == (1400, 860)]
    assert len(scatter) == 3
    # With [0.20, 0.40, 0.80], the middle diamond must sit on the 0.40 tick.
    assert scatter[0].getpixel((727, 486)) == (91, 155, 213)
    assert scatter[0].getpixel((300, 301)) == (255, 192, 0)  # 0.68 reference line.
    summary = next(chart for chart in charts if chart.size == (1400, 800))
    has_orange = (237, 125, 49) in {color for _, color in summary.getcolors(summary.width * summary.height)}
    assert has_orange == (mode == "survey")
print("DOCX typography, package, numbering, chart scale and privacy checks passed")
