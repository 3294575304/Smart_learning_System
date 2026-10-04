import importlib.util
import json
import os
import sys
import tempfile
import zipfile
from pathlib import Path
from docx import Document

sys.dont_write_bytecode = True

spec = importlib.util.spec_from_file_location("custom", "services/quality-reports/custom-template.py")
custom = importlib.util.module_from_spec(spec)
spec.loader.exec_module(custom)
root = Path(".data/tmp").resolve()
root.mkdir(parents=True, exist_ok=True)
with tempfile.TemporaryDirectory(dir=root) as temp:
    template, output = Path(temp) / "template.docx", Path(temp) / "output.docx"
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run("课程：{{课程").bold = True
    paragraph.add_run("名称}} / {{平均分}}")
    doc.sections[0].header.paragraphs[0].text = "{{学期}}"
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "成绩分析"
    table.cell(1, 0).text = "签名"
    doc.save(template)
    parts, roots = custom.load(template)
    slots, targets = custom.discover(roots)
    assert len(slots) == 5, slots
    fields = {"courseName": "示例课程", "mean": "78.50", "term": "2026-2", "gradeAnalysis": "第一行\n第二行"}
    mappings = []
    for slot in slots:
        field = {"课程名称": "courseName", "平均分": "mean", "学期": "term"}.get(slot["label"], "gradeAnalysis" if slot["label"].startswith("成绩分析") else "BLANK")
        mappings.append({"slotId": slot["id"], "field": field})
    payload = Path(temp) / "payload.json"
    payload.write_text(json.dumps({"fields": fields, "mappings": mappings}), encoding="utf-8")
    import sys
    sys.argv = ["custom", "render", str(template), str(output), str(payload)]
    custom.run()
    actual = Document(output)
    assert actual.paragraphs[0].text == "课程：示例课程 / 78.50"
    assert actual.paragraphs[0].runs[0].bold is True
    assert actual.sections[0].header.paragraphs[0].text == "2026-2"
    assert actual.tables[0].cell(0, 1).text == "第一行\n第二行"
    assert actual.tables[0].cell(1, 1).text == ""
    with zipfile.ZipFile(output) as result:
        assert result.read("word/styles.xml") == parts["word/styles.xml"]
    # Security cases must reject the entire upload.
    bad = Path(temp) / "bad.docx"
    for name, data in [("../evil", b"x"), ("word/vbaProject.bin", b"x"), ("word/embeddings/object.bin", b"x"), ("word/_rels/evil.rels", b'<Relationships><Relationship TargetMode="External"/></Relationships>'), ("word/evil.xml", b'<!DOCTYPE x [<!ENTITY x "bad">]><x/>'), ("word/big.bin", b"x" * (11 * 1024 * 1024))]:
        with zipfile.ZipFile(bad, "w", zipfile.ZIP_DEFLATED) as result:
            for key, value in parts.items():
                result.writestr(key, value)
            result.writestr(name, data)
        try:
            custom.load(bad)
            raise AssertionError("Accepted unsafe template: " + name)
        except ValueError:
            pass
print("custom template checks passed")
