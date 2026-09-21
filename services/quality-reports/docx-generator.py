import json
import math
import os
import re
import sys
import zipfile
from copy import deepcopy

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont
from lxml import etree


payload = json.load(open(sys.argv[1], encoding="utf-8"))
template_path = sys.argv[2]
output_path = sys.argv[3]
work_dir = os.path.dirname(output_path)
doc = Document(template_path)
source = payload["source"]
course = source["course"]
classroom = source["classroom"]
statistics = payload["statistics"]
narrative = payload["narrative"]
reviewed = bool(payload.get("reviewed", False))

# Labels belong to this frozen report only; never change source relationship codes.
outcome_numbers = {
    outcome["code"].upper(): index + 1
    for index, outcome in enumerate(source.get("outcomes", []))
}


def display_number(code):
    return str(outcome_numbers.get(str(code).upper(), code))


def display_text(text):
    return re.sub(
        r"(?:课程目标\s*|目标\s*)?(?<![A-Za-z0-9])OBJ\s*[-‐‑–—－]\s*(\d+)(?![A-Za-z0-9])",
        lambda match: f"课程目标{outcome_numbers.get('OBJ-' + match[1], int(match[1]))}",
        str(text),
        flags=re.IGNORECASE,
    )


def font_path():
    candidates = [
        r"C:\Windows\Fonts\simsun.ttc",
        "/usr/share/fonts/truetype/msttcorefonts/simsun.ttc",
    ]
    return next((path for path in candidates if os.path.exists(path)), None)


def chart_font(size, bold=False, latin=False):
    candidates = [
        r"C:\Windows\Fonts\timesbd.ttf" if bold else r"C:\Windows\Fonts\times.ttf",
        "/usr/share/fonts/truetype/msttcorefonts/Times_New_Roman.ttf",
    ] if latin else [font_path()]
    path = next((item for item in candidates if item and os.path.exists(item)), None)
    if not path:
        raise RuntimeError("QUALITY_REPORT_TEMPLATE_FONT_UNAVAILABLE")
    return ImageFont.truetype(path, size)


def apply_run_font(run, size=10.5, bold=False, color=None, name="宋体"):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    properties = run._element.get_or_add_rPr()
    fonts = properties.get_or_add_rFonts()
    for theme in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme", "csTheme"):
        fonts.attrib.pop(qn(f"w:{theme}"), None)
    fonts.set(qn("w:eastAsia"), name)
    fonts.set(qn("w:cs"), "Times New Roman")
    size_cs = properties.find(qn("w:szCs"))
    if size_cs is None:
        size_cs = OxmlElement("w:szCs")
        properties.append(size_cs)
    size_cs.set(qn("w:val"), str(round(size * 2)))


def format_paragraph(paragraph, center=False, before=0, after=0, line=1.5, keep_next=False, indent=False):
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    paragraph.paragraph_format.keep_with_next = keep_next
    paragraph.paragraph_format.first_line_indent = Pt(21 if indent else 0)
    paragraph.paragraph_format.keep_together = False


def clear_cell(cell):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    return cell.paragraphs[0]


def set_cell(cell, text, bold=False, center=False, size=10.5):
    paragraph = clear_cell(cell)
    format_paragraph(paragraph, center=center, line=1.2)
    run = paragraph.add_run(display_text(text))
    apply_run_font(run, size=size, bold=bold)
    return paragraph


def add_text(paragraph, text, bold=False, size=10.5, color=None):
    run = paragraph.add_run(display_text(text))
    apply_run_font(run, size=size, bold=bold, color=color)
    return run


def shade_cell(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def repeat_header(row):
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def dont_split(row):
    properties = row._tr.get_or_add_trPr()
    properties.append(OxmlElement("w:cantSplit"))


def set_cell_width(cell, width):
    cell.width = Inches(width)
    properties = cell._tc.get_or_add_tcPr()
    width_element = properties.find(qn("w:tcW"))
    if width_element is None:
        width_element = OxmlElement("w:tcW")
        properties.append(width_element)
    width_element.set(qn("w:w"), str(round(width * 1440)))
    width_element.set(qn("w:type"), "dxa")


def set_fixed_layout(table, widths):
    table.autofit = False
    properties = table._tbl.tblPr
    layout = properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        properties.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid_columns = table._tbl.tblGrid.gridCol_lst
    for grid_column, width in zip(grid_columns, widths):
        grid_column.set(qn("w:w"), str(round(width * 1440)))
    for row in table.rows:
        handled = set()
        for index, cell in enumerate(row.cells):
            identity = id(cell._tc)
            if identity in handled:
                continue
            handled.add(identity)
            span_element = cell._tc.get_or_add_tcPr().find(qn("w:gridSpan"))
            span = int(span_element.get(qn("w:val"))) if span_element is not None else 1
            set_cell_width(cell, sum(widths[index : index + span]))


def style_table(table, header_rows=1, body_size=10.5, widths=None):
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    if widths:
        set_fixed_layout(table, widths)
    for row_index, row in enumerate(table.rows):
        dont_split(row)
        if row_index < header_rows:
            repeat_header(row)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade_cell(cell, "FFFFFF")
            for paragraph in cell.paragraphs:
                format_paragraph(paragraph, center=True, line=1.25, before=3, after=3)
                for run in paragraph.runs:
                    apply_run_font(
                        run,
                        size=run.font.size.pt if run.font.size else body_size,
                        bold=bool(run.bold),
                    )


def chart_canvas(title, width=1400, height=800):
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    if title:
        draw.text((width / 2, 20), title, anchor="ma", fill="black", font=chart_font(34, True))
    return image, draw, width, height


def grouped_outcome_chart(items, output):
    image, draw, width, height = chart_canvas("")
    draw.rectangle((1, 1, width - 2, height - 2), outline="#D9D9D9", width=2)
    left, top, right, bottom = 120, 70, width - 55, height - 155
    draw.line((left, bottom, right, bottom), fill="#A6A6A6", width=2)
    for tick in range(6):
        y = bottom - (bottom - top) * tick / 5
        draw.line((left, y, right, y), fill="#D9D9D9", width=2)
        draw.text((left - 15, y), f"{tick / 5:.2f}", anchor="rm", fill="#595959", font=chart_font(32, latin=True))
    has_survey = any(item.get("surveyNormalized") is not None for item in items)
    gap = (right - left) / max(len(items), 1)
    for index, item in enumerate(items):
        bar_width = gap * (0.22 if has_survey else 0.30)
        pair_width = bar_width * (2 if has_survey else 1)
        x = left + index * gap + (gap - pair_width) / 2
        values = [(item["attainmentIndex"], "#5B9BD5")]
        if item.get("surveyNormalized") is not None:
            values.append((item["surveyNormalized"], "#ED7D31"))
        for offset, (value, color) in enumerate(values):
            normalized = max(0.0, min(1.0, float(value)))
            x0 = x + offset * bar_width
            y = bottom - normalized * (bottom - top)
            draw.rectangle((x0, y, x0 + bar_width - 3, bottom), fill=color)
            draw.text((x0 + bar_width / 2, y - 10), f"{normalized:.2f}", anchor="ms", fill="#595959", font=chart_font(34, latin=True))
        draw.text((left + (index + 0.5) * gap, bottom + 24), f"目标{display_number(item['code'])}", anchor="ma", fill="#595959", font=chart_font(32))
    legend_x = width / 2 - (330 if has_survey else 130)
    legend_y = height - 48
    draw.rectangle((legend_x, legend_y - 12, legend_x + 24, legend_y + 12), fill="#5B9BD5")
    draw.text((legend_x + 36, legend_y), "达成度计算值", anchor="lm", fill="#595959", font=chart_font(34))
    if has_survey:
        draw.rectangle((legend_x + 330, legend_y - 12, legend_x + 354, legend_y + 12), fill="#ED7D31")
        draw.text((legend_x + 366, legend_y), "学生问卷调查值", anchor="lm", fill="#595959", font=chart_font(34))
    image.save(output)


def normalize_score(value):
    number = float(value)
    return number / 100 if abs(number) > 1.5 else number


def scatter_chart(code, scores, threshold, output):
    values = [normalize_score(value) for value in scores]
    expected = normalize_score(threshold)
    image, draw, width, height = chart_canvas(f"目标{display_number(code)}达成情况", height=860)
    left, top, right, bottom = 120, 90, width - 65, height - 110
    # The points, grid and threshold share one scale; 0.80 must sit on 0.80.
    def y_for(value):
        return bottom - max(0.0, min(1.0, value)) * (bottom - top)
    for tick in range(6):
        y = y_for(tick / 5)
        draw.line((left, y, right, y), fill="#A6A6A6", width=2)
        draw.text((left - 15, y), f"{tick / 5:.2f}", anchor="rm", fill="black", font=chart_font(30, latin=True))
    draw.line((left, top, left, bottom, right, bottom), fill="#A6A6A6", width=2)
    step = max(1, math.ceil((len(values) - 1) / 5))
    tick_indices = sorted(set([1, len(values), *range(1, len(values) + 1, step)]))
    def x_for(index):
        return left + (right - left) * index / max(len(values) - 1, 1)
    for number in tick_indices:
        x = x_for(number - 1)
        draw.line((x, bottom, x, bottom - 10), fill="#A6A6A6", width=2)
        draw.text((x, bottom + 20), str(number), anchor="ma", fill="black", font=chart_font(30, latin=True))
    radius = 10 if len(values) <= 40 else 7
    for index, value in enumerate(values):
        x, y = x_for(index), y_for(value)
        draw.polygon(((x, y - radius), (x + radius, y), (x, y + radius), (x - radius, y)), fill="#5B9BD5")
    draw.line((left, y_for(expected), right, y_for(expected)), fill="#FFC000", width=4)
    draw.text((right, 55), f"期望值 {expected:.2f}", anchor="ra", fill="#806000", font=chart_font(28))
    image.save(output)


def add_picture(cell, path, width=5.1, heading=None):
    paragraph = cell.add_paragraph()
    format_paragraph(paragraph, center=True, line=1.0)
    paragraph.paragraph_format.keep_together = True
    if heading:
        add_text(paragraph, heading, bold=True)
        paragraph.add_run().add_break()
    paragraph.add_run().add_picture(path, width=Inches(width))


def add_caption(cell, text, bold=False):
    # One paragraph keeps an image and its caption together without Word's
    # keepNext promoting an entire enclosing table row to the next page.
    paragraph = cell.paragraphs[-1]
    paragraph.add_run().add_break()
    add_text(paragraph, text, size=9, bold=bold)


def objective_label(title):
    compact = str(title or "").strip()
    if compact.endswith("目标") and len(compact) <= 8:
        return compact[:-2]
    return compact if compact and len(compact) <= 8 else "课程目标"


def reported_attainment(outcome):
    computed = outcome.get("computedAttainmentIndex")
    return computed if computed is not None else outcome.get("attainmentIndex")


def add_multiline_text(cell, text, size=10.5):
    lines = [line.strip() for line in str(text).splitlines() if line.strip()]
    for line in lines:
        paragraph = cell.add_paragraph()
        format_paragraph(paragraph, line=1.5)
        add_text(paragraph, line, size=size)


def cover_field(paragraph_index, value):
    paragraph = doc.paragraphs[paragraph_index]
    if len(paragraph.runs) < 2:
        return
    paragraph.runs[1].text = f"  {value or ' '}  "


def cover_cell(cell, value):
    paragraph = cell.paragraphs[0]
    original_run = next((run for run in paragraph.runs if run.text), None)
    properties = deepcopy(original_run._r.rPr) if original_run is not None else None
    paragraph.clear()
    run = paragraph.add_run(value)
    if properties is not None:
        run._r.append(properties)
    else:
        apply_run_font(run, size=18, bold=True)


def first_year(*values):
    for value in values:
        match = re.search(r"20\d{2}", str(value or ""))
        if match:
            return match.group(0)
    return ""


def restore_template_parts(reference_path, generated_path):
    preserve_parts = {
        "_rels/.rels",
        "customXml/_rels/item1.xml.rels",
        "customXml/_rels/item2.xml.rels",
        "word/_rels/fontTable.xml.rels",
        "word/_rels/header1.xml.rels",
        "word/header1.xml",
        "word/footer1.xml",
        "word/styles.xml",
    }
    patched_path = generated_path + ".patched"
    with zipfile.ZipFile(reference_path, "r") as reference, zipfile.ZipFile(
        generated_path, "r"
    ) as generated, zipfile.ZipFile(
        patched_path, "w", compression=zipfile.ZIP_DEFLATED
    ) as patched:
        for info in generated.infolist():
            data = (
                reference.read(info.filename)
                if info.filename in preserve_parts
                else generated.read(info.filename)
            )
            if info.filename == "word/fontTable.xml":
                # The supplied ideal report declares this title face with a
                # YaHei alternate and no embedded override. The blank template
                # embeds a different title face, so equal rFonts alone differs
                # visibly in Word on the same machine.
                fonts = etree.fromstring(data)
                for font in fonts.findall(qn("w:font")):
                    if font.get(qn("w:name")) != "方正小标宋简体":
                        continue
                    for embedded in font.findall(qn("w:embedRegular")):
                        font.remove(embedded)
                    alternate = font.find(qn("w:altName"))
                    if alternate is None:
                        alternate = OxmlElement("w:altName")
                        font.insert(0, alternate)
                    alternate.set(qn("w:val"), "微软雅黑")
                data = etree.tostring(fonts, xml_declaration=True, encoding="UTF-8", standalone=True)
            patched.writestr(info, data)
    os.replace(patched_path, generated_path)


# Cover slots are edited in place so the original underlines and paragraph tabs remain.
cover_cell(doc.tables[0].cell(0, 0), course["name"])
cover_cell(doc.tables[3].cell(0, 0), course["teacherName"])
cover_field(9, course.get("college", ""))
cover_field(10, course.get("major", ""))
cover_field(11, first_year(course.get("majorClass"), course.get("term")))
cover_field(12, classroom.get("name", ""))

# Course metadata table.
metadata = doc.tables[4]
set_cell(metadata.cell(0, 1), course["name"], bold=True, center=True, size=12)
set_cell(metadata.cell(0, 4), course["courseNature"], bold=True, center=True, size=12)
set_cell(metadata.cell(1, 1), course["term"], bold=True, center=True, size=12)
set_cell(metadata.cell(1, 4), f"{course['credits']:g}", bold=True, center=True, size=12)
set_cell(metadata.cell(2, 2), course.get("majorClass") or classroom["name"], bold=True, center=True, size=12)
set_cell(metadata.cell(2, 4), len(source["students"]), bold=True, center=True, size=12)
set_cell(metadata.cell(3, 2), course["teacherName"], bold=True, center=True, size=12)

# Course objectives.
main = doc.tables[5]
objective_cell = main.cell(1, 0)
paragraph = clear_cell(objective_cell)
for index, outcome in enumerate(source.get("outcomes", [])):
    current = paragraph if index == 0 else objective_cell.add_paragraph()
    format_paragraph(current, line=1.25)
    add_text(
        current,
        f"{index + 1}.{objective_label(outcome.get('title'))}：",
        bold=True,
    )
    description = outcome.get("description")
    add_text(current, description or outcome["title"])
if not source.get("outcomes"):
    add_text(paragraph, "当前数据源未提供正式课程目标。")

# Grade composition and distribution table. The 2024 reference keeps this area
# compact; omitting the redundant bar chart preserves the intended pagination.
grade_cell = main.cell(3, 0)
paragraph = clear_cell(grade_cell)
format_paragraph(paragraph, line=1.5, before=7.8, after=7.8, keep_next=True)
add_text(paragraph, "1.课程总评成绩构成", bold=True)
add_text(paragraph, "（考核方式及所占权重，与教学大纲保持一致）")
paragraph = grade_cell.add_paragraph()
format_paragraph(paragraph, line=1.5)
add_text(paragraph, narrative["gradeComposition"])
paragraph = grade_cell.add_paragraph()
format_paragraph(paragraph, before=7.8, after=7.8, line=1.5, keep_next=True)
add_text(paragraph, "2.总评成绩分布", bold=True)
distribution = statistics["distribution"]
table = grade_cell.add_table(rows=3, cols=len(distribution) + 1)
headers = ["分段"] + [item["label"] for item in distribution]
counts = ["人数"] + [str(item["count"]) for item in distribution]
ratios = ["比例"] + [f"{item['ratio'] * 100:.2f}%" for item in distribution]
for row, values in zip(table.rows, (headers, counts, ratios)):
    for cell, value in zip(row.cells, values):
        set_cell(cell, value, center=True, size=10.5)
style_table(table, header_rows=1, widths=[0.85] + [1.07] * len(distribution))
paragraph = grade_cell.add_paragraph()
format_paragraph(paragraph, before=3, line=1.5, indent=True)
add_text(paragraph, narrative["gradeAnalysis"])

# Outcome table, summary chart, AI draft analysis and per-outcome evidence chart.
outcome_cell = main.cell(5, 0)
paragraph = clear_cell(outcome_cell)
format_paragraph(paragraph, line=1.5, before=7.8, after=7.8, keep_next=True)
add_text(paragraph, "1.课程目标达成情况——定量统计", bold=True)
add_text(paragraph, "（评分标准与课程大纲保持一致）")
paragraph = outcome_cell.add_paragraph()
format_paragraph(paragraph, line=1.2, keep_next=True)
add_text(paragraph, "注：A 为各考核方式平均分按课程目标占比加权后的合计，B 为对应满分加权合计，达成度=A/B；结果保留2位小数。", size=9)
outcomes = statistics.get("outcomes", [])
components = source.get("components", [])
if outcomes:
    component_count = len(components)
    column_count = component_count + 5
    outcome_table = outcome_cell.add_table(rows=len(outcomes) + 4, cols=column_count)
    component_start = 2
    component_end = component_start + component_count - 1
    average_column = component_end + 1
    maximum_column = component_end + 2
    ratio_column = component_end + 3
    set_cell(outcome_table.cell(0, 0).merge(outcome_table.cell(3, 0)), "课程目标", bold=True, center=True, size=8)
    set_cell(
        outcome_table.cell(0, 1).merge(outcome_table.cell(0, component_end)),
        "考核方式及所占权重",
        bold=True,
        center=True,
        size=8,
    )
    for column, label in (
        (average_column, "加权平均分\nA"),
        (maximum_column, "加权总分\nB"),
        (ratio_column, "达成度\nA/B"),
    ):
        set_cell(
            outcome_table.cell(0, column).merge(outcome_table.cell(3, column)),
            label,
            bold=True,
            center=True,
            size=8,
        )
    set_cell(outcome_table.cell(1, 1), "考核方式", bold=True, center=True, size=8)
    set_cell(outcome_table.cell(2, 1), "总分", bold=True, center=True, size=8)
    set_cell(outcome_table.cell(3, 1), "平均分", bold=True, center=True, size=8)
    component_means = {item["code"]: item["mean"] for item in statistics.get("componentMeans", [])}
    for offset, component in enumerate(components):
        column = component_start + offset
        set_cell(outcome_table.cell(1, column), component["name"], bold=True, center=True, size=8)
        set_cell(outcome_table.cell(2, column), "100", center=True, size=8)
        mean = component_means.get(component["code"])
        set_cell(outcome_table.cell(3, column), "—" if mean is None else f"{mean:.2f}", center=True, size=8)
    for row_index, outcome in enumerate(outcomes, start=4):
        allocations = {item["componentCode"]: item["allocationRate"] for item in outcome.get("componentAllocations", [])}
        set_cell(outcome_table.cell(row_index, 0), display_number(outcome["code"]), center=True, size=10.5)
        set_cell(outcome_table.cell(row_index, 1), "权重", center=True, size=8)
        for offset, component in enumerate(components):
            allocation = allocations.get(component["code"])
            set_cell(
                outcome_table.cell(row_index, component_start + offset),
                "—" if allocation is None else f"{allocation * 100:.1f}%",
                center=True,
                size=8,
            )
        set_cell(
            outcome_table.cell(row_index, average_column),
            "—" if outcome.get("weightedAverage") is None else f"{outcome['weightedAverage']:.2f}",
            center=True,
            size=8,
        )
        set_cell(
            outcome_table.cell(row_index, maximum_column),
            "—" if outcome.get("weightedMaximum") is None else f"{outcome['weightedMaximum']:.2f}",
            center=True,
            size=8,
        )
        set_cell(
            outcome_table.cell(row_index, ratio_column),
            "—" if reported_attainment(outcome) is None else f"{reported_attainment(outcome):.2f}",
            center=True,
            size=8,
        )
    target_widths = [22.6 / 72, 35.15 / 72] + [280 / 72 / max(component_count, 1)] * component_count + [39.7 / 72, 39.7 / 72, 39.75 / 72]
    # The reference uses 10.5 pt labels and 12 pt numerical values, white cells.
    for row_index, row in enumerate(outcome_table.rows):
        row.height = Pt([20.95, 62.75, 30.5, 30.5][row_index] if row_index < 4 else 37.75)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for column, cell in enumerate(row.cells):
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    numeric = row_index >= 2 and column >= component_start and bool(re.fullmatch(r"[\d.%—]+", run.text))
                    apply_run_font(run, size=12 if numeric else 10.5, bold=bool(run.bold) and not numeric)
    style_table(outcome_table, header_rows=2, widths=target_widths)
paragraph = outcome_cell.add_paragraph()
format_paragraph(paragraph, before=5, line=1.2, keep_next=True)
add_text(paragraph, "2.课程目标达成情况——定性分析", bold=True)
chartable_outcomes = [
    outcome
    for outcome in outcomes
    if reported_attainment(outcome) is not None
]
if chartable_outcomes:
    chartable_outcomes = [
        {**outcome, "attainmentIndex": reported_attainment(outcome)}
        for outcome in chartable_outcomes
    ]
    summary_chart = os.path.join(work_dir, "outcome-summary.png")
    grouped_outcome_chart(chartable_outcomes, summary_chart)
    add_picture(outcome_cell, summary_chart, 4.61)
    comparison_available = any(
        outcome.get("surveyNormalized") is not None for outcome in chartable_outcomes
    )
    add_caption(
        outcome_cell,
        "图1 课程目标达成度与学生自评对照图"
        if comparison_available
        else "图1 课程目标达成情况分析图（未纳入学生自评）",
    )
paragraph = outcome_cell.add_paragraph()
format_paragraph(paragraph, line=1.5, indent=True)
add_text(paragraph, narrative["outcomeAnalysis"])
paragraph = outcome_cell.add_paragraph()
format_paragraph(paragraph, before=3, line=1.3, keep_next=True)
add_text(paragraph, "3.课程目标个体评价分布图及分析", bold=True)
paragraph = outcome_cell.add_paragraph()
format_paragraph(paragraph, line=1.25)
add_text(paragraph, "下图展示学生各课程目标的达成情况，横坐标为学生序号，纵坐标为达成度，横线表示课程目标期望值。")
details = {item["code"]: item["analysis"] for item in narrative.get("outcomeDetails", [])}
for index, outcome in enumerate(outcomes, start=2):
    heading = f"课程目标{display_number(outcome['code'])}达成情况分析："
    if outcome.get("studentScores") and outcome.get("threshold") is not None:
        path = os.path.join(work_dir, f"outcome-{index}.png")
        scatter_chart(outcome["code"], outcome["studentScores"], outcome["threshold"], path)
        # Word can ignore keepNext within a split outer table row. Keep the
        # heading, picture and caption in one indivisible paragraph instead.
        add_picture(outcome_cell, path, 4.0, heading=heading)
        add_caption(outcome_cell, f"图{index} 课程目标{display_number(outcome['code'])}达成情况分布", bold=True)
    else:
        paragraph = outcome_cell.add_paragraph()
        format_paragraph(paragraph, before=3, line=1.25)
        add_text(paragraph, heading, bold=True)
    paragraph = outcome_cell.add_paragraph()
    format_paragraph(paragraph, line=1.5, indent=True)
    add_text(paragraph, details.get(outcome["code"], "暂无分析文字。"))

# Questionnaire aggregate, AI course summary and improvement measures.
evaluation = doc.tables[6]
paragraph = set_cell(evaluation.cell(1, 0), narrative["studentEvaluation"])
format_paragraph(paragraph, line=1.5, indent=True)
summary_cell = evaluation.cell(5, 0)
paragraph = clear_cell(summary_cell)
format_paragraph(paragraph, line=1.35, keep_next=True)
add_text(paragraph, "1.课程总结：", bold=True)
add_multiline_text(summary_cell, narrative["courseSummary"])
paragraph = summary_cell.add_paragraph()
format_paragraph(paragraph, before=4, line=1.35, keep_next=True)
add_text(paragraph, "2.持续改进措施：", bold=True)
add_multiline_text(summary_cell, narrative["improvementMeasures"])
survey_title = source.get("survey", {}).get("title") if source.get("survey") else course["name"]
set_cell(evaluation.cell(7, 0), f"附件3.1  《{survey_title}》学生调查问卷\n……", size=10.5)
for paragraph in evaluation.cell(6, 0).paragraphs:
    paragraph.paragraph_format.keep_with_next = True
dont_split(evaluation.rows[6])
dont_split(evaluation.rows[7])

# The template's review opinions, signatures and dates are intentionally untouched.
# Keep the review form together; it starts a new page whenever it does not fit.
# Unlike an unconditional break, this does not isolate a short appendix page.
for paragraph in doc.paragraphs:
    if paragraph.text.strip() == "四、审核意见":
        paragraph.paragraph_format.keep_with_next = True
for row in doc.tables[7].rows:
    dont_split(row)
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.keep_with_next = True
doc.tables[7].rows[1].height = Pt(175.6)
for paragraph in doc.paragraphs[23:25]:
    paragraph.paragraph_format.keep_with_next = True

# Filled narrative rows must flow like the reference. A reserved blank-form
# minimum is reapplied after a split in Word; keepNext inside a short outer row
# can move the whole row, producing half-empty pages even when text fits.
for table_index, row_index in ((5, 1), (5, 3), (5, 5), (6, 1), (6, 5), (6, 7)):
    row = doc.tables[table_index].rows[row_index]
    row.height = None
    for paragraph in row.cells[0].paragraphs:
        paragraph.paragraph_format.keep_with_next = False

# Attach each appendix to the preceding narrative instead of creating a page
# containing only the appendix. The following section uses pageBreakBefore,
# avoiding a standalone blank break paragraph after the table.
for table_index in (5, 6):
    table = doc.tables[table_index]
    for paragraph in table.cell(5, 0).paragraphs[-2:]:
        paragraph.paragraph_format.keep_with_next = True
    for row_index in (6, 7):
        table.rows[row_index].height = None
        dont_split(table.rows[row_index])
        for paragraph in table.cell(row_index, 0).paragraphs:
            paragraph.paragraph_format.keep_with_next = row_index == 6
for paragraph in doc.paragraphs:
    if paragraph.text.strip() == "三、课程质量综合评价与持续改进":
        paragraph.paragraph_format.page_break_before = True
        previous = paragraph._p.getprevious()
        if previous is not None and previous.tag == qn("w:p") and not previous.xpath(".//w:t"):
            for line_break in previous.xpath(".//w:br[@w:type='page']"):
                line_break.getparent().remove(line_break)
            # Keep the empty source paragraph, bookmarks and package identity.
            properties = previous.get_or_add_pPr()
            spacing = properties.get_or_add_spacing()
            spacing.set(qn("w:line"), "1")
            spacing.set(qn("w:lineRule"), "exact")
doc.core_properties.title = f"{course['name']}课程教学质量分析报告"
doc.core_properties.comments = (
    "AI 生成分析初稿，已经授课教师审核确认；课程评价小组、学院审核意见、签字和日期留空。"
    if reviewed
    else "AI 生成分析初稿，须经授课教师审核确认后方可作为正式报告；课程评价小组、学院审核意见、签字和日期留空。"
)
settings = doc.settings._element
update_fields = settings.find(qn("w:updateFields"))
if update_fields is None:
    update_fields = OxmlElement("w:updateFields")
    settings.append(update_fields)
update_fields.set(qn("w:val"), "true")
doc.save(output_path)
restore_template_parts(template_path, output_path)
